from collections.abc import Generator
from datetime import date, time
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.core.database import Base, get_db
from app.main import app
from app.modules.auth.models import AuditLog, Permission, User
from app.modules.auth.seed import seed_initial_admin


@pytest.fixture
def db_session() -> Generator[Session, None, None]:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    with factory() as session:
        yield session
    engine.dispose()


@pytest.fixture
def client(db_session: Session) -> Generator[TestClient, None, None]:
    def override_get_db() -> Generator[Session, None, None]:
        yield db_session

    app.dependency_overrides[get_db] = override_get_db
    try:
        with TestClient(app) as test_client:
            yield test_client
    finally:
        app.dependency_overrides.clear()


@pytest.fixture
def admin(db_session: Session) -> tuple[User, str]:
    seeded = seed_initial_admin(db_session)
    assert seeded.password is not None
    actor = db_session.scalar(select(User).where(User.username == seeded.username))
    assert actor is not None
    db_session.commit()
    return actor, seeded.password


def _login(client: TestClient, actor: User, password: str) -> dict[str, str]:
    response = client.post(
        "/api/v1/auth/login",
        json={"username": actor.username, "password": password},
    )
    assert response.status_code == 200
    return {"Authorization": "Bearer " + response.json()["access_token"]}


def _create_employee(client: TestClient, headers: dict[str, str], code: str) -> int:
    response = client.post(
        "/api/v1/employees",
        headers=headers,
        json={
            "code": code,
            "name": f"Employee {code}",
            "mobile": "01712345678",
            "join_date": "2025-01-01",
            "base_salary": "10000.00",
        },
    )
    assert response.status_code == 201
    return response.json()["id"]


def _attendance_workbook(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Employee Code", "Attendance Date", "Status", "In Time", "Out Time"])
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_attendance_create_update_date_filter_and_audit(
    client: TestClient,
    db_session: Session,
    admin: tuple[User, str],
) -> None:
    actor, password = admin
    headers = _login(client, actor, password)
    employee_id = _create_employee(client, headers, "EMP-001")

    created = client.post(
        "/api/v1/attendance",
        headers=headers,
        json={
            "employee_id": employee_id,
            "att_date": "2026-10-16",
            "status": "PRESENT",
            "in_time": "08:30",
            "out_time": "17:00",
        },
    )
    assert created.status_code == 201
    assert created.json()["employee"]["code"] == "EMP-001"
    assert created.json()["in_time"] == "08:30:00"

    duplicate = client.post(
        "/api/v1/attendance",
        headers=headers,
        json={
            "employee_id": employee_id,
            "att_date": "2026-10-16",
            "status": "LATE",
        },
    )
    assert duplicate.status_code == 409

    updated = client.patch(
        f"/api/v1/attendance/{created.json()['id']}",
        headers=headers,
        json={"status": "LATE", "in_time": None},
    )
    assert updated.status_code == 200
    assert updated.json()["status"] == "LATE"
    assert updated.json()["in_time"] is None

    listed = client.get(
        "/api/v1/attendance",
        headers=headers,
        params={"from_date": "2026-10-16", "to_date": "2026-10-16"},
    )
    assert listed.status_code == 200
    assert len(listed.json()) == 1
    assert listed.json()[0]["status"] == "LATE"
    assert client.get(
        "/api/v1/attendance",
        headers=headers,
        params={"from_date": "2026-10-17", "to_date": "2026-10-16"},
    ).status_code == 422

    audits = list(
        db_session.scalars(
            select(AuditLog).where(AuditLog.entity_type == "attendance")
        )
    )
    assert [entry.action for entry in audits] == ["CREATE", "UPDATE"]
    assert audits[-1].before_json["status"] == "PRESENT"
    assert audits[-1].after_json["status"] == "LATE"


def test_attendance_import_creates_valid_rows_and_reports_rejected_rows(
    client: TestClient,
    admin: tuple[User, str],
) -> None:
    actor, password = admin
    headers = _login(client, actor, password)
    first_employee = _create_employee(client, headers, "EMP-001")
    second_employee = _create_employee(client, headers, "EMP-002")
    existing = client.post(
        "/api/v1/attendance",
        headers=headers,
        json={
            "employee_id": second_employee,
            "att_date": "2026-10-16",
            "status": "PRESENT",
        },
    )
    assert existing.status_code == 201

    result = client.post(
        "/api/v1/attendance/import",
        content=_attendance_workbook(
            [
                ["EMP-001", "2026-10-16", "present", "08:30", "17:00"],
                ["EMP-002", "2026-10-16", "ABSENT", None, None],
                ["UNKNOWN", "2026-10-16", "PRESENT", None, None],
                ["EMP-001", "2026-10-17", "HOLIDAY", None, None],
            ]
        ),
        headers={
            **headers,
            "Content-Type": (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
        },
    )
    assert result.status_code == 200
    assert result.json()["created"] == 1
    assert result.json()["rejected"] == 3
    assert [item["row_number"] for item in result.json()["errors"]] == [3, 4, 5]
    assert first_employee > 0


def test_attendance_workbook_template_export_and_invalid_upload(
    client: TestClient,
    admin: tuple[User, str],
) -> None:
    actor, password = admin
    headers = _login(client, actor, password)
    employee_id = _create_employee(client, headers, "EMP-001")
    created = client.post(
        "/api/v1/attendance",
        headers=headers,
        json={
            "employee_id": employee_id,
            "att_date": "2026-10-16",
            "status": "PRESENT",
            "in_time": "08:30",
            "out_time": "17:00",
        },
    )
    assert created.status_code == 201

    template = client.get("/api/v1/attendance/template", headers=headers)
    assert template.status_code == 200
    assert template.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    template_book = load_workbook(BytesIO(template.content), read_only=True)
    assert next(template_book.active.iter_rows(values_only=True)) == (
        "Employee Code",
        "Attendance Date",
        "Status",
        "In Time",
        "Out Time",
    )
    template_book.close()

    exported = client.get(
        "/api/v1/attendance/export",
        headers=headers,
        params={"from_date": "2026-10-16", "to_date": "2026-10-16"},
    )
    assert exported.status_code == 200
    export_book = load_workbook(BytesIO(exported.content), read_only=True)
    rows = list(export_book.active.iter_rows(values_only=True))
    assert rows[1][0] == "EMP-001"
    assert rows[1][1].date() == date(2026, 10, 16)
    assert rows[1][2:] == ("PRESENT", time(8, 30), time(17, 0))
    export_book.close()

    invalid = client.post(
        "/api/v1/attendance/import",
        headers={
            **headers,
            "Content-Type": (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
        },
        content=b"not an xlsx workbook",
    )
    assert invalid.status_code == 422


def test_attendance_validation_and_permission_denial(
    client: TestClient,
    db_session: Session,
    admin: tuple[User, str],
) -> None:
    actor, password = admin
    headers = _login(client, actor, password)
    employee_id = _create_employee(client, headers, "EMP-001")

    invalid = client.post(
        "/api/v1/attendance",
        headers=headers,
        json={
            "employee_id": employee_id,
            "att_date": "2026-10-16",
            "status": "HOLIDAY",
            "in_time": "25:00",
        },
    )
    assert invalid.status_code == 422

    permission = db_session.scalar(
        select(Permission).where(Permission.code == "employees.attendance.read")
    )
    assert permission is not None
    actor.roles[0].permissions.remove(permission)
    db_session.commit()
    headers = _login(client, actor, password)
    assert client.get(
        "/api/v1/attendance/options", headers=headers
    ).status_code == 403
    assert client.get(
        "/api/v1/attendance",
        headers=headers,
        params={"from_date": "2026-10-16", "to_date": "2026-10-16"},
    ).status_code == 403
