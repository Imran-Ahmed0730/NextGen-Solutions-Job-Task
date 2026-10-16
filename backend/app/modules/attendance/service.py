from collections.abc import Sequence
from datetime import date, datetime, time
from io import BytesIO
from zipfile import BadZipFile

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import TypeAdapter, ValidationError
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.modules.attendance import repository
from app.modules.attendance.models import Attendance
from app.modules.attendance.schemas import AttendanceImportLine
from app.modules.auth.audit import record_audit
from app.modules.auth.models import User
from app.modules.employees.models import Employee

ATTENDANCE_HEADERS = (
    "Employee Code",
    "Attendance Date",
    "Status",
    "In Time",
    "Out Time",
)


def list_attendance_for_range(
    db: Session, from_date: date, to_date: date
) -> list[Attendance]:
    if from_date > to_date:
        raise HTTPException(status_code=422, detail="Start date cannot be after end date")
    return repository.list_attendance_for_range(db, from_date, to_date)


def list_employee_options(db: Session) -> list[Employee]:
    return repository.list_employee_options(db)


def create_attendance(
    db: Session, data: dict[str, object], actor: User, ip: str | None = None
) -> Attendance:
    employee_id = data.get("employee_id")
    att_date = data.get("att_date")
    if not isinstance(employee_id, int) or employee_id <= 0:
        raise HTTPException(status_code=422, detail="Choose an employee")
    if not isinstance(att_date, date):
        raise HTTPException(status_code=422, detail="Attendance date is required")
    employee = repository.get_employee(db, employee_id)
    if employee is None or employee.deleted_at is not None:
        raise HTTPException(status_code=404, detail="Employee not found")
    if repository.get_employee_attendance(db, employee_id, att_date) is not None:
        raise _duplicate_attendance_error()

    attendance = Attendance(
        employee_id=employee_id,
        att_date=att_date,
        status=_status(data.get("status")),
    )
    setattr(attendance, "in_time", _time(data.get("in_time"), "In time"))
    setattr(attendance, "out_time", _time(data.get("out_time"), "Out time"))
    db.add(attendance)
    try:
        db.flush()
    except IntegrityError as exc:
        db.rollback()
        raise _duplicate_attendance_error() from exc
    record_audit(
        db,
        actor_id=actor.id,
        entity_type="attendance",
        entity_id=attendance.id,
        action="CREATE",
        after=_snapshot(attendance),
        ip=ip,
    )
    db.commit()
    result = repository.get_attendance(db, attendance.id)
    if result is None:
        raise RuntimeError("Created attendance record could not be reloaded")
    return result


def update_attendance(
    db: Session,
    attendance_id: int,
    data: dict[str, object],
    actor: User,
    ip: str | None = None,
) -> Attendance:
    attendance = repository.get_attendance(db, attendance_id, for_update=True)
    if attendance is None:
        raise HTTPException(status_code=404, detail="Attendance record not found")
    before = _snapshot(attendance)
    if "status" in data:
        attendance.status = _status(data["status"])
    if "in_time" in data:
        setattr(attendance, "in_time", _time(data["in_time"], "In time"))
    if "out_time" in data:
        setattr(attendance, "out_time", _time(data["out_time"], "Out time"))

    record_audit(
        db,
        actor_id=actor.id,
        entity_type="attendance",
        entity_id=attendance.id,
        action="UPDATE",
        before=before,
        after=_snapshot(attendance),
        ip=ip,
    )
    db.commit()
    result = repository.get_attendance(db, attendance.id)
    if result is None:
        raise RuntimeError("Updated attendance record could not be reloaded")
    return result


def import_attendance(
    db: Session,
    rows: list[dict[str, object]],
    actor: User,
    ip: str | None = None,
) -> dict[str, object]:
    created = 0
    errors: list[dict[str, object]] = []
    seen: set[tuple[int, date]] = set()
    adapter: TypeAdapter[AttendanceImportLine] = TypeAdapter(AttendanceImportLine)

    for index, row in enumerate(rows, start=2):
        row_number = row.get("row_number")
        if not isinstance(row_number, int) or row_number < 2:
            row_number = index
        raw_line = dict(row)
        raw_line.pop("row_number", None)
        raw_status = raw_line.get("status")
        if isinstance(raw_status, str):
            raw_line["status"] = raw_status.strip().upper()
        try:
            line = adapter.validate_python(raw_line)
        except ValidationError as exc:
            errors.append(
                {"row_number": row_number, "message": _validation_message(exc)}
            )
            continue

        employee = repository.get_employee_by_code(db, line.employee_code.strip())
        if employee is None or employee.deleted_at is not None:
            errors.append({"row_number": row_number, "message": "Employee code not found"})
            continue
        key = (employee.id, line.att_date)
        if key in seen or repository.get_employee_attendance(
            db, employee.id, line.att_date
        ):
            errors.append(
                {
                    "row_number": row_number,
                    "message": "Attendance already exists for this employee and date",
                }
            )
            continue
        seen.add(key)

        attendance = Attendance(
            employee_id=employee.id,
            att_date=line.att_date,
            status=line.status,
        )
        setattr(attendance, "in_time", line.in_time)
        setattr(attendance, "out_time", line.out_time)
        db.add(attendance)
        try:
            db.flush()
        except IntegrityError:
            db.rollback()
            errors.append(
                {
                    "row_number": row_number,
                    "message": "Attendance already exists for this employee and date",
                }
            )
            continue
        record_audit(
            db,
            actor_id=actor.id,
            entity_type="attendance",
            entity_id=attendance.id,
            action="CREATE",
            after=_snapshot(attendance),
            ip=ip,
        )
        db.commit()
        created += 1

    return {"created": created, "rejected": len(errors), "errors": errors}


def parse_attendance_workbook(data: bytes) -> list[dict[str, object]]:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
        raise HTTPException(
            status_code=422,
            detail="The uploaded file is not a valid .xlsx workbook",
        ) from exc

    try:
        if not workbook.worksheets:
            raise HTTPException(status_code=422, detail="The workbook has no worksheet")
        sheet = workbook.worksheets[0]
        first_row = next(sheet.iter_rows(min_row=1, max_row=1), ())
        headers = {
            str(cell.value).strip().casefold(): index
            for index, cell in enumerate(first_row, start=1)
            if cell.value is not None
        }
        required = {header.casefold() for header in ATTENDANCE_HEADERS}
        if not required.issubset(headers):
            raise HTTPException(
                status_code=422,
                detail=f"Workbook must contain these columns: {', '.join(ATTENDANCE_HEADERS)}",
            )

        indexes = {header: headers[header.casefold()] for header in ATTENDANCE_HEADERS}
        rows: list[dict[str, object]] = []
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
            line: dict[str, object] = {
                "row_number": row_number,
                "employee_code": _excel_value(cells, indexes["Employee Code"], "text"),
                "att_date": _excel_value(cells, indexes["Attendance Date"], "date"),
                "status": _excel_value(cells, indexes["Status"], "text"),
                "in_time": _excel_value(cells, indexes["In Time"], "time"),
                "out_time": _excel_value(cells, indexes["Out Time"], "time"),
            }
            if not any(
                value not in (None, "")
                for key, value in line.items()
                if key != "row_number"
            ):
                continue
            status_value = line["status"]
            if isinstance(status_value, str):
                line["status"] = status_value.strip().upper()
            rows.append(line)
            if len(rows) > 1000:
                raise HTTPException(
                    status_code=422,
                    detail="Import no more than 1,000 attendance rows at a time",
                )
        if not rows:
            raise HTTPException(status_code=422, detail="The workbook contains no attendance rows")
        return rows
    finally:
        workbook.close()


def build_attendance_workbook(records: list[Attendance] | None = None) -> bytes:
    workbook = Workbook()
    sheet = workbook.worksheets[0]
    sheet.title = "Attendance"
    sheet.append(ATTENDANCE_HEADERS)
    sheet.freeze_panes = "A2"
    for column, width in enumerate((22, 18, 16, 14, 14), start=1):
        sheet.column_dimensions[chr(64 + column)].width = width

    for record in records or []:
        sheet.append(
            [
                record.employee.code,
                record.att_date,
                record.status,
                record.in_time,
                record.out_time,
            ]
        )
        sheet.cell(sheet.max_row, 2).number_format = "yyyy-mm-dd"
        sheet.cell(sheet.max_row, 4).number_format = "hh:mm"
        sheet.cell(sheet.max_row, 5).number_format = "hh:mm"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _status(value: object) -> str:
    if not isinstance(value, str) or value not in {"PRESENT", "ABSENT", "LEAVE", "LATE"}:
        raise HTTPException(status_code=422, detail="Choose a valid attendance status")
    return str(value)


def _time(value: object, label: str) -> time | None:
    if value is None:
        return None
    if not isinstance(value, time):
        raise HTTPException(status_code=422, detail=f"{label} is invalid")
    return value


def _duplicate_attendance_error() -> HTTPException:
    return HTTPException(
        status_code=409,
        detail="Attendance has already been recorded for this employee and date",
    )


def _snapshot(attendance: Attendance) -> dict[str, object]:
    return {
        "employee_id": attendance.employee_id,
        "att_date": attendance.att_date.isoformat(),
        "status": attendance.status,
        "in_time": attendance.in_time.isoformat() if attendance.in_time else None,
        "out_time": attendance.out_time.isoformat() if attendance.out_time else None,
    }


def _validation_message(error: ValidationError) -> str:
    return "; ".join(
        f"{'.'.join(str(part) for part in item['loc'])}: {item['msg']}"
        for item in error.errors()
    )


def _excel_value(cells: Sequence[Cell | MergedCell], column: int, kind: str) -> object:
    value = cells[column - 1].value if column <= len(cells) else None
    if value is None:
        return None
    if kind == "text":
        return str(value).strip()
    if kind == "date":
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        return value
    if kind == "time":
        if isinstance(value, datetime):
            value = value.time()
        if isinstance(value, time):
            return value.isoformat(timespec="minutes")
        return str(value).strip() or None
    return value
